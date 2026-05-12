#!/usr/bin/env python3
"""PA Channel Export — PA 渠道数据导出器"""

import argparse, csv, sys
from pathlib import Path


def detect_encoding(filepath):
    for enc in ["utf-8-sig", "utf-8", "gbk", "latin-1"]:
        try:
            with open(filepath, "r", encoding=enc) as f: f.read(4096)
            return enc
        except (UnicodeDecodeError, UnicodeError): continue
    return "utf-8"


# Key fraud-related columns to highlight
HIGHLIGHT_COLUMNS = [
    "Blocked Reason",
    "blocked_reason",
    "Rejected Reason",
    "rejected_reason",
    "Fraud Type",
    "fraud_type",
    "Risk Score",
    "risk_score",
    "Protection Status",
    "protection_status",
]


def export_pa_data(
    input_path, output_path,
    channel=None,
    app_id=None,
    highlight_columns=None,
    dry_run=False,
):
    """
    Export PA anti-fraud data to styled Excel.
    
    Highlights key fraud reason columns with distinctive formatting.
    Preserves all original fields from the source CSV.
    """
    encoding = detect_encoding(input_path)
    highlight_cols = set(HIGHLIGHT_COLUMNS + (highlight_columns or []))

    rows = []
    fieldnames = []
    matched = 0

    with open(input_path, "r", encoding=encoding, newline="") as infile:
        reader = csv.DictReader(infile)
        fieldnames = reader.fieldnames or []
        
        for row in reader:
            if channel:
                ch_val = row.get("Channel", "") or row.get("channel", "")
                if channel.lower() not in ch_val.lower():
                    continue
            if app_id:
                aid_val = row.get("App ID", "") or row.get("app_id", "")
                if app_id.lower() not in aid_val.lower():
                    continue
            
            matched += 1
            rows.append(row)

    if dry_run:
        return {"matched_rows": matched, "total_fields": len(fieldnames),
                "highlight_cols": [c for c in fieldnames if c.lower() in [h.lower() for h in highlight_cols]],
                "output": "(dry run)"}

    # Try writing Excel with openpyxl styling
    try:
        _write_styled_excel(output_path, rows, fieldnames, highlight_cols)
    except ImportError:
        # Fallback to CSV if openpyxl not available
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8", newline="") as out:
            writer = csv.DictWriter(out, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    highlighted = [c for c in fieldnames if c.lower() in [h.lower() for h in highlight_cols]]
    return {"matched_rows": matched, "total_fields": len(fieldnames),
            "highlight_cols": highlighted, "output": output_path}


def _write_styled_excel(output_path, rows, fieldnames, highlight_cols):
    """Write styled Excel with highlighted fraud columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "PA Export"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill_orange = PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid")
    header_fill_highlight = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
    cell_fill_yellow = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    # Write headers
    for col_idx, fname in enumerate(fieldnames, 1):
        cell = ws.cell(row=1, column=col_idx, value=fname)
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin_border
        
        if fname in highlight_cols or fname.lower() in [h.lower() for h in highlight_cols]:
            cell.fill = header_fill_highlight
        else:
            cell.fill = header_fill_orange

    # Write data rows
    for row_idx, row in enumerate(rows, 2):
        for col_idx, fname in enumerate(fieldnames, 1):
            val = row.get(fname, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if fname in highlight_cols or fname.lower() in [h.lower() for h in highlight_cols]:
                cell.fill = cell_fill_yellow

    # Auto-adjust column widths (rough heuristic)
    for col_idx, fname in enumerate(fieldnames, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            len(str(fname)) + 2, 12
        )

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="PA Anti-Fraud Data Export")
    parser.add_argument("-i", "--input", required=True, help="Input PA CSV file")
    parser.add_argument("-o", "--output", required=True, help="Output Excel/CSV path")
    parser.add_argument("--channel", help="Filter by channel name")
    parser.add_argument("--app-id", help="Filter by App ID")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not Path(args.input).exists():
        print(f"Not found: {args.input}", file=sys.stderr); sys.exit(1)

    result = export_pa_data(args.input, args.output, args.channel,
                            args.app_id, dry_run=args.dry_run)
    
    print(f"\nPA Export Summary")
    print(f"{'='*40}")
    print(f"  Matched rows:     {result['matched_rows']:,}")
    print(f"  Total fields:    {result['total_fields']}")
    print(f"  Highlight cols:  {result['highlight_cols']}")
    print(f"  Output:          {result['output']}")


if __name__ == "__main__":
    main()
